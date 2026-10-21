"""Лабораторная работа №1: консольный генератор экзаменационных билетов.

Запрашивает фамилию и имя студента, выдаёт случайный номер билета 1..20
и сразу дописывает запись в journal.xlsx.
"""
import os
import random
import re
import sys
import unicodedata
from datetime import datetime
from typing import Callable, Protocol

from openpyxl import Workbook, load_workbook

JOURNAL_FILE = "journal.xlsx"
HEADER = ["Last name", "First name", "Номер билета", "Дата и время"]
COLUMN_WIDTHS = {"A": 20, "B": 20, "C": 15, "D": 22}
TICKET_MIN, TICKET_MAX = 1, 20
ESC = "\x1b"


class ExitRequested(Exception):
    """Пользователь нажал ESC."""


class NameValidator(Protocol):
    def validate(self, raw: str) -> str:
        """Возвращает нормализованное имя или сообщает об ошибке."""


class TicketNumberGenerator(Protocol):
    def generate(self) -> int:
        """Возвращает номер билета."""


class RecordRepository(Protocol):
    def append(self, last_name: str, first_name: str, ticket: int) -> None:
        """Сохраняет запись в журнал."""


class UserInterface(Protocol):
    def ask_name(self, prompt: str) -> str:
        """Запрашивает и возвращает проверенное имя."""

    def show_message(self, message: str) -> None:
        """Выводит сообщение пользователю."""

    def show_ticket(self, ticket: int) -> None:
        """Показывает выданный билет."""

    def wait_for_retry(self) -> None:
        """Ждёт подтверждения перед повторной попыткой сохранения."""


class UnicodeNameValidator:
    """Проверяет имя или фамилию на буквы, заглавную первую букву и разделители."""

    _separators = re.compile(r"['’ -]")

    def validate(self, raw: str) -> str:
        value = unicodedata.normalize("NFC", raw.strip())
        if not value:
            raise ValueError("Поле не может быть пустым.")
        if not value[0].isupper():
            raise ValueError("Имя и фамилия должны начинаться с заглавной буквы.")
        parts = self._separators.split(value)
        if not all(part and part.isalpha() for part in parts):
            raise ValueError(
                "Допустимы только буквы, пробелы, дефис и апостроф."
            )
        return value


def clean_name(raw):
    """Обрезает пробелы по краям. Возвращает None, если строка пустая."""
    value = raw.strip()
    return value or None


def generate_ticket(rng=random):
    return rng.randint(TICKET_MIN, TICKET_MAX)


class RandomTicketGenerator:
    def __init__(self, rng=random):
        self._rng = rng

    def generate(self) -> int:
        return generate_ticket(self._rng)


def append_record(path, last_name, first_name, ticket, when=None):
    """Дописывает одну строку в конец журнала и сразу сохраняет файл.

    Если файла нет, создаёт его с шапкой. Существующие строки не меняются.
    PermissionError пробрасывается наверх (файл открыт в Excel).
    """
    when = when or datetime.now()
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Journal"
        ws.append(HEADER)
        for column, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[column].width = width
    ws.append([last_name, first_name, ticket, when.strftime("%d.%m.%Y %H:%M:%S")])
    wb.save(path)


class ExcelJournalRepository:
    def __init__(self, path=JOURNAL_FILE):
        self._path = path

    def append(self, last_name: str, first_name: str, ticket: int) -> None:
        append_record(self._path, last_name, first_name, ticket)


# ---------- консольный ввод ----------
def _read_key():
    """Считывает одну клавишу без ожидания Enter."""
    if os.name == "nt":
        import msvcrt
        ch = msvcrt.getwch()
        if ch in ("\x00", "\xe0"):  # стрелки, F-клавиши: пропускаем второй код
            msvcrt.getwch()
            return ""
        return ch
    if not sys.stdin.isatty():
        ch = sys.stdin.read(1)
        if ch == "":
            raise EOFError
        return ch
    import termios
    import tty
    fd = sys.stdin.fileno()
    old = termios.tcgetattr(fd)
    try:
        tty.setraw(fd)
        return sys.stdin.read(1)
    finally:
        termios.tcsetattr(fd, termios.TCSADRAIN, old)


def ask(prompt, validator=None):
    """Запрашивает допустимое имя. ESC в начале ввода -> ExitRequested."""
    if validator is None:
        validator = UnicodeNameValidator()
    while True:
        print(prompt, end="", flush=True)
        characters = []
        while True:
            key = _read_key()
            if key == ESC and not characters:
                print()
                raise ExitRequested
            if key == "\x03":  # Ctrl+C
                raise KeyboardInterrupt
            if key in ("\r", "\n"):
                print()
                break
            if key in ("\x08", "\x7f"):
                if characters:
                    characters.pop()
                    print("\b \b", end="", flush=True)
                continue
            if not key or key in ("\x00", "\xe0"):
                continue
            characters.append(key)
            print(key, end="", flush=True)

        value = clean_name("".join(characters))
        try:
            return validator.validate(value or "")
        except ValueError as error:
            print(f"{error} Повторите ввод.")


class ConsoleUserInterface:
    def __init__(self, name_validator: NameValidator):
        self._name_validator = name_validator

    def ask_name(self, prompt: str) -> str:
        return ask(prompt, self._name_validator)

    def show_message(self, message: str) -> None:
        print(message)

    def show_ticket(self, ticket: int) -> None:
        print(f"Билет № {ticket}")

    def wait_for_retry(self) -> None:
        input("Закройте файл и нажмите Enter, чтобы повторить сохранение...")


def save_with_retry(last_name, first_name, ticket, path=JOURNAL_FILE):
    def retry_message():
        print(f"Файл {path} открыт в другой программе (например, в Excel) "
              "и недоступен для записи.")
        input("Закройте файл и нажмите Enter, чтобы повторить сохранение...")

    _retry_on_permission_error(
        lambda: append_record(path, last_name, first_name, ticket),
        retry_message,
    )


def _retry_on_permission_error(
    action: Callable[[], None], on_error: Callable[[], None]
) -> None:
    while True:
        try:
            action()
            return
        except PermissionError:
            on_error()


class TicketApplication:
    def __init__(
        self,
        ticket_generator: TicketNumberGenerator,
        repository: RecordRepository,
        user_interface: UserInterface,
    ):
        self._ticket_generator = ticket_generator
        self._repository = repository
        self._user_interface = user_interface

    def run(self) -> None:
        self._user_interface.show_message("Для выхода нажмите ESC")
        try:
            self._run_until_exit()
        except (ExitRequested, KeyboardInterrupt, EOFError):
            self._user_interface.show_message("Работа завершена.")

    def _run_until_exit(self) -> None:
        while True:
            last_name = self._user_interface.ask_name("Last name: ")
            first_name = self._user_interface.ask_name("First name: ")
            ticket = self._ticket_generator.generate()
            self._user_interface.show_ticket(ticket)
            self._save_with_retry(last_name, first_name, ticket)
            self._user_interface.show_message("")

    def _save_with_retry(
        self, last_name: str, first_name: str, ticket: int
    ) -> None:
        def show_retry_message() -> None:
            self._user_interface.show_message(
                "Файл журнала открыт в другой программе (например, в Excel) "
                "и недоступен для записи."
            )
            self._user_interface.wait_for_retry()

        _retry_on_permission_error(
            lambda: self._repository.append(last_name, first_name, ticket),
            show_retry_message,
        )


def main():
    name_validator = UnicodeNameValidator()
    application = TicketApplication(
        ticket_generator=RandomTicketGenerator(),
        repository=ExcelJournalRepository(),
        user_interface=ConsoleUserInterface(name_validator),
    )
    application.run()


if __name__ == "__main__":
    main()
