import random
import sys
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import ticket_generator as tg  # noqa: E402


def rows(path):
    return [list(r) for r in load_workbook(path).active.iter_rows(values_only=True)]


def test_clean_name_trims_spaces():
    assert tg.clean_name("  Anghis  ") == "Anghis"


@pytest.mark.parametrize("raw", ["", "   ", "\t"])
def test_clean_name_rejects_empty(raw):
    assert tg.clean_name(raw) is None


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("  Дмитрий  ", "Дмитрий"),
        ("Jean-Luc", "Jean-Luc"),
        ("O’Connor", "O’Connor"),
        ("Van der Meer", "Van der Meer"),
    ],
)
def test_name_validator_accepts_letters_and_name_separators(raw, expected):
    assert tg.UnicodeNameValidator().validate(raw) == expected


@pytest.mark.parametrize(
    "raw",
    [
        "", "   ", "Иван2", "Иван!", "-Иван", "Иван--Петров",
        "иван", "van der Meer",
    ],
)
def test_name_validator_rejects_invalid_names(raw):
    with pytest.raises(ValueError):
        tg.UnicodeNameValidator().validate(raw)


def test_ask_retries_after_invalid_name(monkeypatch, capsys):
    keys = iter(
        ["и", "в", "а", "н", "\r", "И", "в", "а", "н", "\r"]
    )
    monkeypatch.setattr(tg, "_read_key", lambda: next(keys))

    assert tg.ask("Name: ") == "Иван"
    assert "заглавной буквы" in capsys.readouterr().out


def test_ask_backspace_can_delete_first_character(monkeypatch):
    keys = iter(["И", "в", "а", "н", "\x08", "\x08", "\x08", "\x08",
                 "П", "е", "т", "р", "\r"])
    monkeypatch.setattr(tg, "_read_key", lambda: next(keys))

    assert tg.ask("Name: ") == "Петр"


def test_application_uses_injected_components():
    class Generator:
        def generate(self):
            return 7

    class Repository:
        records = []

        def append(self, last_name, first_name, ticket):
            self.records.append((last_name, first_name, ticket))

    class UserInterface:
        def __init__(self):
            self.names = iter(["Иванов", "Иван"])
            self.messages = []
            self.tickets = []

        def ask_name(self, prompt):
            try:
                return next(self.names)
            except StopIteration:
                raise tg.ExitRequested

        def show_message(self, message):
            self.messages.append(message)

        def show_ticket(self, ticket):
            self.tickets.append(ticket)

        def wait_for_retry(self):
            pass

    repository = Repository()
    user_interface = UserInterface()
    application = tg.TicketApplication(
        Generator(), repository, user_interface
    )

    application.run()

    assert repository.records == [("Иванов", "Иван", 7)]
    assert user_interface.tickets == [7]
    assert user_interface.messages[-1] == "Работа завершена."


def test_ticket_in_range():
    rng = random.Random(1)
    assert all(1 <= tg.generate_ticket(rng) <= 20 for _ in range(1000))


def test_creates_file_with_header(tmp_path):
    path = tmp_path / "journal.xlsx"
    tg.append_record(path, "Anghis", "Dmitrii", 7, datetime(2026, 10, 19, 12, 0, 0))
    assert rows(path) == [tg.HEADER, ["Anghis", "Dmitrii", 7, "19.10.2026 12:00:00"]]


def test_appends_without_overwriting(tmp_path):
    path = tmp_path / "journal.xlsx"
    for i in range(3):
        tg.append_record(path, f"L{i}", f"F{i}", i + 1)
    before = rows(path)
    tg.append_record(path, "L3", "F3", 4)
    tg.append_record(path, "L4", "F4", 5)
    after = rows(path)
    assert len(after) == 6
    assert after[:4] == before


def test_locked_file_is_retried(tmp_path, monkeypatch, capsys):
    calls = {"n": 0}

    def flaky(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            raise PermissionError
    monkeypatch.setattr(tg, "append_record", flaky)
    monkeypatch.setattr("builtins.input", lambda *_: "")
    tg.save_with_retry("A", "B", 1, path=str(tmp_path / "j.xlsx"))
    assert calls["n"] == 2
    assert "недоступен для записи" in capsys.readouterr().out
