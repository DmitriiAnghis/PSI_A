"""Создаёт тестовые data/students.xlsx и data/tickets.docx."""
import os

from docx import Document
from openpyxl import Workbook

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")

GROUPS = {
    "TI-235": [("Anghis", "Dmitrii"), ("Popescu", "Ion"), ("Rusu", "Maria"), ("Ciobanu", "Andrei")],
    "TI-236": [("Lungu", "Elena"), ("Moraru", "Victor"), ("Cebotari", "Ana")],
    "TI-237": [],  # пустая группа для проверки
}

TOPICS = [
    "Жизненный цикл ПО", "Требования к ПО", "UML-диаграммы", "Тестирование", "Git и ветвление",
    "Паттерны проектирования", "Архитектура клиент-сервер", "Базы данных", "REST API", "Agile и Scrum",
]


def make_students(path):
    wb = Workbook()
    wb.remove(wb.active)
    for group, students in GROUPS.items():
        ws = wb.create_sheet(group)
        ws.append(["Фамилия", "Имя"])
        for row in students:
            ws.append(list(row))
    wb.save(path)


def make_tickets(path):
    doc = Document()
    for n, topic in enumerate(TOPICS, 1):
        doc.add_paragraph(f"Билет {n}")
        doc.add_paragraph(f"1. Основные понятия темы «{topic}».")
        doc.add_paragraph(f"2. Практический пример применения: «{topic}».")
        doc.add_paragraph(f"3. Преимущества и недостатки подхода «{topic}».")
        doc.add_paragraph("")
    # некорректный билет: только два вопроса, должен быть пропущен
    doc.add_paragraph("Билет 11")
    doc.add_paragraph("1. Вопрос без пары.")
    doc.add_paragraph("2. Ещё один вопрос.")
    doc.save(path)


if __name__ == "__main__":
    os.makedirs(DATA_DIR, exist_ok=True)
    make_students(os.path.join(DATA_DIR, "students.xlsx"))
    make_tickets(os.path.join(DATA_DIR, "tickets.docx"))
    print("Готово:", DATA_DIR)
