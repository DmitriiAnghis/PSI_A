"""Работа с файлами Лаб 2: students.xlsx, tickets.docx, results.xlsx."""
import logging
import os
import random
import re
from datetime import datetime

from docx import Document
from openpyxl import Workbook, load_workbook

log = logging.getLogger("tickets")

RESULTS_HEADER = ["Группа", "Фамилия", "Имя", "Номер билета", "Дата и время", "Повтор (да/нет)"]
RESULTS_WIDTHS = {"A": 12, "B": 18, "C": 16, "D": 15, "E": 22, "F": 17}
QUESTIONS_PER_TICKET = 3

_TICKET_RE = re.compile(r"^\s*Билет\s*№?\s*(\d+)\s*$", re.IGNORECASE)
_QUESTION_RE = re.compile(r"^\s*\d+\s*[.)]\s*(.+?)\s*$")


class SourceError(Exception):
    """Входной файл отсутствует или не читается."""


def _cell_text(value):
    return "" if value is None else str(value).strip()


def load_groups(path):
    """Возвращает {группа: [(фамилия, имя), ...]} в порядке листов students.xlsx."""
    if not os.path.exists(path):
        raise SourceError(f"Не найден файл со студентами: {path}")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise SourceError(f"Не удалось прочитать {path}: {exc}") from exc
    groups = {}
    for ws in wb.worksheets:
        students = []
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            last, first = (_cell_text(v) for v in (tuple(row) + (None, None))[:2])
            if last or first:
                students.append((last, first))
        groups[ws.title] = students
    wb.close()
    return groups


def _is_list_paragraph(paragraph):
    """Абзац оформлен автоматической нумерацией Word (номер не входит в текст)."""
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def parse_ticket_lines(lines):
    """Разбирает строки документа в {номер: [вопросы]}.

    lines: список (текст, это_пункт_нумерованного_списка).
    Билеты, у которых меньше трёх вопросов, пропускаются с записью в лог.
    """
    tickets = {}
    current, questions = None, []

    def flush():
        if current is None:
            return
        if len(questions) < QUESTIONS_PER_TICKET:
            log.warning("Билет %s пропущен: найдено вопросов %d из %d",
                        current, len(questions), QUESTIONS_PER_TICKET)
        elif current in tickets:
            log.warning("Билет %s пропущен: номер повторяется", current)
        else:
            tickets[current] = questions[:QUESTIONS_PER_TICKET]

    for text, is_list_item in lines:
        text = text.strip()
        if not text:
            continue
        header = _TICKET_RE.match(text)
        if header:
            flush()
            current, questions = int(header.group(1)), []
            continue
        question = _QUESTION_RE.match(text)
        if current is None:
            log.warning("Строка вне билета пропущена: %r", text)
        elif question:
            questions.append(question.group(1))
        elif is_list_item:
            questions.append(text)
        else:
            log.warning("Нераспознанная строка в билете %s: %r", current, text)
    flush()
    return tickets


def load_tickets(path):
    """Читает tickets.docx и возвращает {номер билета: [3 вопроса]}."""
    if not os.path.exists(path):
        raise SourceError(f"Не найден файл с билетами: {path}")
    try:
        doc = Document(path)
    except Exception as exc:
        raise SourceError(f"Не удалось прочитать {path}: {exc}") from exc
    lines = [(p.text, _is_list_paragraph(p)) for p in doc.paragraphs]
    tickets = parse_ticket_lines(lines)
    if not tickets:
        raise SourceError(f"В {path} не найдено ни одного корректного билета")
    return tickets


def _key(group, last, first):
    return (group.strip().casefold(), last.strip().casefold(), first.strip().casefold())


def find_first_ticket(results_path, group, last, first):
    """Номер билета из самой первой записи студента в журнале или None."""
    if not os.path.exists(results_path):
        return None
    wb = load_workbook(results_path, read_only=True)
    try:
        target = _key(group, last, first)
        for row in wb.active.iter_rows(min_row=2, max_col=4, values_only=True):
            g, l, f, number = (tuple(row) + (None,) * 4)[:4]
            if number is not None and _key(_cell_text(g), _cell_text(l), _cell_text(f)) == target:
                return int(number)
    finally:
        wb.close()
    return None


def append_result(results_path, group, last, first, ticket, repeat, when=None):
    """Дописывает строку в конец results.xlsx и сразу сохраняет.

    PermissionError пробрасывается наверх (файл открыт в Excel).
    """
    when = when or datetime.now()
    if os.path.exists(results_path):
        wb = load_workbook(results_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(RESULTS_HEADER)
        for column, width in RESULTS_WIDTHS.items():
            ws.column_dimensions[column].width = width
    ws.append([group, last, first, ticket, when.strftime("%d.%m.%Y %H:%M:%S"),
               "да" if repeat else "нет"])
    wb.save(results_path)


def assign_ticket(results_path, tickets, group, last, first, rng=random):
    """Возвращает (номер, повтор): первый номер студента или случайный из tickets."""
    previous = find_first_ticket(results_path, group, last, first)
    if previous is not None:
        return previous, True
    return rng.choice(sorted(tickets)), False


# --- добавление студента ---

_NAME_RE = re.compile(r"^[^\W\d_]+(?:[ '’-][^\W\d_]+)*$")


class InvalidNameError(ValueError):
    """Фамилия или имя не прошли проверку."""


class DuplicateStudentError(ValueError):
    """Такой студент уже есть в группе."""


def validate_name(text, field="Поле"):
    """Возвращает имя без лишних пробелов или бросает InvalidNameError.

    Правила: не пусто, только буквы (допускаются пробел, дефис и апостроф
    между частями), без цифр, первая буква заглавная.
    """
    value = " ".join(_cell_text(text).split())
    if not value:
        raise InvalidNameError(f"{field}: поле не может быть пустым.")
    if any(ch.isdigit() for ch in value):
        raise InvalidNameError(f"{field}: цифры не допускаются.")
    if not _NAME_RE.match(value):
        raise InvalidNameError(
            f"{field}: допускаются только буквы, пробел, дефис и апостроф между частями.")
    if not value[0].isupper():
        raise InvalidNameError(f"{field}: должно начинаться с заглавной буквы.")
    return value


def add_student(path, group, last, first):
    """Проверяет имя и дописывает студента в лист группы students.xlsx.

    Возвращает (фамилия, имя) после очистки. PermissionError пробрасывается
    наверх (файл открыт в Excel).
    """
    last = validate_name(last, "Фамилия")
    first = validate_name(first, "Имя")
    if not os.path.exists(path):
        raise SourceError(f"Не найден файл со студентами: {path}")
    wb = load_workbook(path)
    if group not in wb.sheetnames:
        wb.close()
        raise SourceError(f"В {os.path.basename(path)} нет листа группы {group}")
    ws = wb[group]
    target = (last.casefold(), first.casefold())
    for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        l, f = (_cell_text(v) for v in (tuple(row) + (None, None))[:2])
        if (l.casefold(), f.casefold()) == target:
            wb.close()
            raise DuplicateStudentError(f"{last} {first} уже есть в группе {group}.")
    ws.append([last, first])
    wb.save(path)
    return last, first


# --- группы ---

STUDENTS_HEADER = ["Фамилия", "Имя"]
_GROUP_RE = re.compile(r"^[^\W\d_][\w-]*$")
_SHEET_FORBIDDEN = set('[]:*?/\\')


class DuplicateGroupError(ValueError):
    """Группа с таким названием уже есть."""


def validate_group(text):
    """Название группы: начинается с заглавной буквы, дальше буквы, цифры, дефис.

    Excel ограничивает имя листа 31 символом и запрещает []:*?/\\.
    """
    value = _cell_text(text).replace(" ", "")
    if not value:
        raise InvalidNameError("Группа: поле не может быть пустым.")
    if len(value) > 31:
        raise InvalidNameError("Группа: не длиннее 31 символа.")
    if _SHEET_FORBIDDEN & set(value) or not _GROUP_RE.match(value):
        raise InvalidNameError("Группа: только буквы, цифры и дефис, например TI-238.")
    if not value[0].isupper():
        raise InvalidNameError("Группа: должна начинаться с заглавной буквы.")
    return value


def add_group(path, name):
    """Создаёт в students.xlsx лист новой группы с шапкой. Возвращает название."""
    name = validate_group(name)
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    if name.casefold() in (title.casefold() for title in wb.sheetnames):
        wb.close()
        raise DuplicateGroupError(f"Группа {name} уже есть.")
    ws = wb.create_sheet(name)
    ws.append(STUDENTS_HEADER)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    wb.save(path)
    return name


def remove_student(path, group, last, first):
    """Удаляет первую строку студента из листа группы. Возвращает True, если нашёл."""
    wb = load_workbook(path)
    if group not in wb.sheetnames:
        wb.close()
        raise SourceError(f"В {os.path.basename(path)} нет листа группы {group}")
    ws = wb[group]
    target = (last.strip().casefold(), first.strip().casefold())
    for index, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
        l, f = (_cell_text(v) for v in (tuple(row) + (None, None))[:2])
        if (l.casefold(), f.casefold()) == target:
            ws.delete_rows(index)
            wb.save(path)
            return True
    wb.close()
    return False


# --- журнал ---

def load_results(path):
    """Строки results.xlsx без шапки: [(группа, фамилия, имя, билет, дата, повтор)]."""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True)
    try:
        rows = []
        for row in wb.active.iter_rows(min_row=2, max_col=6, values_only=True):
            values = (tuple(row) + (None,) * 6)[:6]
            if any(v is not None for v in values):
                rows.append(tuple(_cell_text(v) for v in values))
        return rows
    finally:
        wb.close()


def results_stats(rows):
    """Сводка по журналу: выдано, повторов, студентов, самый частый билет."""
    issued = len(rows)
    repeats = sum(1 for r in rows if r[5].casefold() == "да")
    students = len({_key(r[0], r[1], r[2]) for r in rows})
    counts = {}
    for r in rows:
        if r[5].casefold() != "да" and r[3]:
            counts[r[3]] = counts.get(r[3], 0) + 1
    top = max(counts.items(), key=lambda kv: (kv[1], -int(kv[0]) if kv[0].isdigit() else 0),
              default=None)
    return {"issued": issued, "repeats": repeats, "students": students,
            "top_ticket": top[0] if top else None, "top_count": top[1] if top else 0}


# --- редактирование билетов в tickets.docx ---

class TicketError(ValueError):
    """Ошибка в номере или вопросах билета."""


def validate_questions(questions):
    """Ровно три непустых вопроса, каждый не короче 3 символов."""
    cleaned = [" ".join(_cell_text(q).split()) for q in questions]
    if len(cleaned) != QUESTIONS_PER_TICKET:
        raise TicketError(f"В билете должно быть {QUESTIONS_PER_TICKET} вопроса.")
    for i, question in enumerate(cleaned, 1):
        if not question:
            raise TicketError(f"Вопрос {i}: поле не может быть пустым.")
        if len(question) < 3:
            raise TicketError(f"Вопрос {i}: слишком короткий.")
    if len({q.casefold() for q in cleaned}) != len(cleaned):
        raise TicketError("Вопросы в одном билете не должны повторяться.")
    return cleaned


def _ticket_blocks(doc):
    """{номер: [абзацы билета]} — заголовок «Билет N» и всё до следующего заголовка."""
    blocks, current = {}, None
    for paragraph in doc.paragraphs:
        header = _TICKET_RE.match(paragraph.text)
        if header:
            current = int(header.group(1))
            blocks.setdefault(current, []).append(paragraph)
        elif current is not None:
            blocks[current].append(paragraph)
    return blocks


def _open_tickets(path):
    if not os.path.exists(path):
        raise SourceError(f"Не найден файл с билетами: {path}")
    return Document(path)


def _remove_paragraphs(paragraphs):
    for paragraph in paragraphs:
        element = paragraph._element
        element.getparent().remove(element)


def _ticket_paragraph_texts(number, questions):
    return [f"Билет {number}"] + [f"{i}. {q}" for i, q in enumerate(questions, 1)] + [""]


def add_ticket(path, questions, number=None):
    """Дописывает билет в конец tickets.docx. Номер по умолчанию — следующий свободный."""
    questions = validate_questions(questions)
    doc = _open_tickets(path)
    used = set(_ticket_blocks(doc))
    if number is None:
        number = max(used, default=0) + 1
    else:
        try:
            number = int(str(number).strip())
        except ValueError:
            raise TicketError("Номер билета должен быть целым числом.") from None
        if number < 1:
            raise TicketError("Номер билета должен быть больше нуля.")
        if number in used:
            raise TicketError(f"Билет № {number} уже есть.")
    for text in _ticket_paragraph_texts(number, questions):
        doc.add_paragraph(text)
    doc.save(path)
    return number


def update_ticket(path, number, questions):
    """Заменяет вопросы билета, сохраняя его место в документе."""
    questions = validate_questions(questions)
    doc = _open_tickets(path)
    blocks = _ticket_blocks(doc)
    if number not in blocks:
        raise TicketError(f"Билета № {number} нет в tickets.docx.")
    header, *rest = blocks[number]
    _remove_paragraphs(rest)
    anchor = header
    for text in _ticket_paragraph_texts(number, questions)[1:]:
        new = doc.add_paragraph(text)  # создаём в конце и переносим после заголовка
        anchor._element.addnext(new._element)
        anchor = new
    doc.save(path)
    return questions


def remove_ticket(path, number):
    """Удаляет билет (заголовок и его вопросы). Возвращает True, если нашёл."""
    doc = _open_tickets(path)
    blocks = _ticket_blocks(doc)
    if number not in blocks:
        return False
    _remove_paragraphs(blocks[number])
    doc.save(path)
    return True


# --- от себя: удаление пустой группы, экспорт журнала ---

def remove_group(path, name):
    """Удаляет лист группы, только если в ней нет студентов."""
    wb = load_workbook(path)
    if name not in wb.sheetnames:
        wb.close()
        raise SourceError(f"Группы {name} нет.")
    ws = wb[name]
    has_students = any(any(_cell_text(v) for v in row)
                       for row in ws.iter_rows(min_row=2, max_col=2, values_only=True))
    if has_students:
        wb.close()
        raise InvalidNameError(f"В группе {name} есть студенты: сначала удалите их.")
    if len(wb.sheetnames) == 1:
        wb.close()
        raise InvalidNameError("Нельзя удалить последнюю группу.")
    wb.remove(ws)
    wb.save(path)


def export_results_csv(rows, path):
    """Сохраняет строки журнала в CSV (разделитель «;», UTF-8 с BOM — открывается в Excel)."""
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(RESULTS_HEADER)
        writer.writerows(rows)
    return path
