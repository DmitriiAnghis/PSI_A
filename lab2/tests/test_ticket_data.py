import random
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import ticket_data as td  # noqa: E402


@pytest.fixture
def students(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    a = wb.create_sheet("TI-235")
    a.append(["Фамилия", "Имя"])
    a.append(["Anghis", "Dmitrii"])
    a.append([" Popescu ", "Ion"])
    b = wb.create_sheet("TI-237")
    b.append(["Фамилия", "Имя"])
    path = tmp_path / "students.xlsx"
    wb.save(path)
    return path


def make_docx(path, paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def rows_of(path, sheet):
    return [list(r) for r in load_workbook(path)[sheet].iter_rows(values_only=True)]


def rows(path):
    return [list(r) for r in load_workbook(path).active.iter_rows(values_only=True)]


# --- студенты по группе ---

def test_groups_are_sheets_and_students_by_group(students):
    groups = td.load_groups(students)
    assert list(groups) == ["TI-235", "TI-237"]
    assert groups["TI-235"] == [("Anghis", "Dmitrii"), ("Popescu", "Ion")]


def test_empty_group_has_no_students(students):
    assert td.load_groups(students)["TI-237"] == []


def test_missing_students_file(tmp_path):
    with pytest.raises(td.SourceError):
        td.load_groups(tmp_path / "nope.xlsx")


# --- парсинг билетов из Word ---

def test_parse_ticket_from_word(tmp_path):
    path = make_docx(tmp_path / "t.docx", [
        "Билет 1", "1. Вопрос А", "2. Вопрос Б", "3. Вопрос В", "",
        "Билет 2", "1. Q1", "2. Q2", "3. Q3",
    ])
    assert td.load_tickets(path) == {1: ["Вопрос А", "Вопрос Б", "Вопрос В"], 2: ["Q1", "Q2", "Q3"]}


def test_ticket_with_less_than_three_questions_is_skipped(tmp_path, caplog):
    path = make_docx(tmp_path / "t.docx", [
        "Билет 1", "1. A", "2. B",
        "Билет 2", "1. A", "2. B", "3. C",
    ])
    assert list(td.load_tickets(path)) == [2]
    assert "Билет 1 пропущен" in caplog.text


def test_unrecognized_header_is_ignored():
    lines = [("Bilet X", False), ("1. A", False), ("Билет 3", False),
             ("1. A", False), ("2. B", False), ("3. C", False)]
    assert td.parse_ticket_lines(lines) == {3: ["A", "B", "C"]}


def test_word_auto_numbered_list_is_parsed():
    lines = [("Билет 1", False), ("A", True), ("B", True), ("C", True)]
    assert td.parse_ticket_lines(lines) == {1: ["A", "B", "C"]}


def test_missing_tickets_file(tmp_path):
    with pytest.raises(td.SourceError):
        td.load_tickets(tmp_path / "nope.docx")


# --- журнал и идемпотентность ---

def test_new_student_gets_random_ticket_and_no_repeat(tmp_path):
    results = tmp_path / "results.xlsx"
    tickets = {n: ["a", "b", "c"] for n in range(1, 11)}
    number, repeat = td.assign_ticket(results, tickets, "TI-235", "Anghis", "Dmitrii", random.Random(0))
    assert number in tickets and repeat is False


def test_repeat_selection_returns_same_ticket(tmp_path):
    results = tmp_path / "results.xlsx"
    tickets = {n: ["a", "b", "c"] for n in range(1, 21)}
    first, _ = td.assign_ticket(results, tickets, "TI-235", "Anghis", "Dmitrii")
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", first, False)
    for _ in range(5):
        number, repeat = td.assign_ticket(results, tickets, "TI-235", "Anghis", "Dmitrii")
        assert (number, repeat) == (first, True)
        td.append_result(results, "TI-235", "Anghis", "Dmitrii", number, repeat)


def test_same_name_in_other_group_is_different_student(tmp_path):
    results = tmp_path / "results.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 7, False)
    assert td.find_first_ticket(results, "TI-236", "Anghis", "Dmitrii") is None


def test_journal_appends_and_keeps_history(tmp_path):
    results = tmp_path / "results.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 7, False)
    td.append_result(results, "TI-235", "Popescu", "Ion", 3, False)
    before = rows(results)
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 7, True)
    after = rows(results)
    assert after[0] == td.RESULTS_HEADER
    assert after[:3] == before
    assert after[3][:4] == ["TI-235", "Anghis", "Dmitrii", 7] and after[3][5] == "да"
    assert [r[5] for r in after[1:]] == ["нет", "нет", "да"]


# --- добавление студента ---

@pytest.mark.parametrize("text, expected", [
    ("Anghis", "Anghis"),
    ("  Дмитрий  ", "Дмитрий"),
    ("Jean-Luc", "Jean-Luc"),
    ("O'Connor", "O'Connor"),
    ("Van  der Meer", "Van der Meer"),
])
def test_validate_name_accepts_correct_names(text, expected):
    assert td.validate_name(text) == expected


@pytest.mark.parametrize("text, message", [
    ("", "пустым"),
    ("   ", "пустым"),
    ("Ivan2", "цифры"),
    ("3", "цифры"),
    ("ivan", "заглавной"),
    ("дмитрий", "заглавной"),
    ("Ivan!", "только буквы"),
    ("-Ivan", "только буквы"),
    ("Ivan--Petrov", "только буквы"),
])
def test_validate_name_rejects_bad_names(text, message):
    with pytest.raises(td.InvalidNameError, match=message):
        td.validate_name(text)


def test_add_student_appends_to_group_sheet(students):
    assert td.add_student(students, "TI-237", " Ciobanu ", "Andrei") == ("Ciobanu", "Andrei")
    groups = td.load_groups(students)
    assert groups["TI-237"] == [("Ciobanu", "Andrei")]
    assert len(groups["TI-235"]) == 2


def test_add_student_rejects_invalid_name_without_writing(students):
    before = students.read_bytes()
    with pytest.raises(td.InvalidNameError):
        td.add_student(students, "TI-235", "popescu", "Ion")
    assert students.read_bytes() == before


def test_add_student_rejects_duplicate_in_same_group(students):
    with pytest.raises(td.DuplicateStudentError):
        td.add_student(students, "TI-235", "Popescu", "Ion")
    td.add_student(students, "TI-237", "Popescu", "Ion")


def test_add_student_unknown_group(students):
    with pytest.raises(td.SourceError):
        td.add_student(students, "TI-999", "Popescu", "Ion")


def test_close_in_excel_is_noop_without_excel(tmp_path):
    from excel_closer import close_in_excel
    assert close_in_excel(tmp_path / "results.xlsx") is False


# --- группы, удаление, журнал ---

@pytest.mark.parametrize("text, expected", [("TI-238", "TI-238"), (" FAF-231 ", "FAF-231"),
                                            ("ТИ-235", "ТИ-235")])
def test_validate_group_accepts(text, expected):
    assert td.validate_group(text) == expected


@pytest.mark.parametrize("text", ["", "ti-238", "238", "TI/238", "TI:1", "-TI", "A" * 32])
def test_validate_group_rejects(text):
    with pytest.raises(td.InvalidNameError):
        td.validate_group(text)


def test_add_group_creates_sheet_with_header(students):
    assert td.add_group(students, "TI-238") == "TI-238"
    groups = td.load_groups(students)
    assert list(groups) == ["TI-235", "TI-237", "TI-238"]
    assert groups["TI-238"] == []
    assert rows_of(students, "TI-238") == [["Фамилия", "Имя"]]
    td.add_student(students, "TI-238", "Rusu", "Maria")
    assert td.load_groups(students)["TI-238"] == [("Rusu", "Maria")]


def test_add_group_rejects_duplicate_case_insensitive(students):
    with pytest.raises(td.DuplicateGroupError):
        td.add_group(students, "TI-235")
    with pytest.raises(td.InvalidNameError):
        td.add_group(students, "ti-235")


def test_remove_student(students):
    assert td.remove_student(students, "TI-235", "Popescu", "Ion") is True
    assert td.load_groups(students)["TI-235"] == [("Anghis", "Dmitrii")]
    assert td.remove_student(students, "TI-235", "Popescu", "Ion") is False


def test_results_and_stats(tmp_path):
    path = tmp_path / "results.xlsx"
    assert td.load_results(path) == []
    td.append_result(path, "TI-235", "Anghis", "Dmitrii", 7, False)
    td.append_result(path, "TI-235", "Popescu", "Ion", 7, False)
    td.append_result(path, "TI-235", "Anghis", "Dmitrii", 7, True)
    td.append_result(path, "TI-236", "Lungu", "Elena", 3, False)
    rows_ = td.load_results(path)
    assert len(rows_) == 4 and rows_[0][:4] == ("TI-235", "Anghis", "Dmitrii", "7")
    stats = td.results_stats(rows_)
    assert stats == {"issued": 4, "repeats": 1, "reassigned": 0, "students": 3, "top_ticket": "7", "top_count": 2}


# --- билеты в tickets.docx, удаление группы, экспорт ---

QS = ["Первый вопрос", "Второй вопрос", "Третий вопрос"]


@pytest.fixture
def tickets_doc(tmp_path):
    return make_docx(tmp_path / "tickets.docx", [
        "Билет 1", "1. Вопрос А", "2. Вопрос Б", "3. Вопрос В", "",
        "Билет 2", "1. Вопрос Г", "2. Вопрос Д", "3. Вопрос Е", "",
    ])


def test_add_ticket_next_number(tickets_doc):
    assert td.add_ticket(tickets_doc, QS) == 3
    tickets = td.load_tickets(tickets_doc)
    assert sorted(tickets) == [1, 2, 3] and tickets[3] == QS


def test_add_ticket_with_number_and_errors(tickets_doc):
    assert td.add_ticket(tickets_doc, QS, number=" 10 ") == 10
    for number in ("2", "abc", "0"):
        with pytest.raises(td.TicketError):
            td.add_ticket(tickets_doc, QS, number=number)


@pytest.mark.parametrize("questions", [["", "Второй", "Третий"], ["Ok", "Второй", "Третий"],
                                       ["Один", "один", "Три"], ["Один", "Два"]])
def test_validate_questions_rejects(questions):
    with pytest.raises(td.TicketError):
        td.validate_questions(questions)


def test_update_ticket_keeps_position(tickets_doc):
    td.update_ticket(tickets_doc, 1, QS)
    tickets = td.load_tickets(tickets_doc)
    assert tickets[1] == QS and tickets[2] == ["Вопрос Г", "Вопрос Д", "Вопрос Е"]
    texts = [p.text for p in Document(tickets_doc).paragraphs if p.text]
    assert texts.index("Билет 1") < texts.index("1. Первый вопрос") < texts.index("Билет 2")
    with pytest.raises(td.TicketError):
        td.update_ticket(tickets_doc, 5, QS)


def test_remove_ticket(tickets_doc):
    assert td.remove_ticket(tickets_doc, 1) is True
    assert list(td.load_tickets(tickets_doc)) == [2]
    assert td.remove_ticket(tickets_doc, 1) is False


def test_remove_group_only_when_empty(students):
    with pytest.raises(td.InvalidNameError):
        td.remove_group(students, "TI-235")
    td.remove_group(students, "TI-237")
    assert list(td.load_groups(students)) == ["TI-235"]
    td.remove_student(students, "TI-235", "Anghis", "Dmitrii")
    td.remove_student(students, "TI-235", "Popescu", "Ion")
    with pytest.raises(td.InvalidNameError, match="последнюю"):
        td.remove_group(students, "TI-235")


def test_export_results_csv(tmp_path):
    out = td.export_results_csv([("TI-235", "Anghis", "Dmitrii", "7", "21.10.2026 10:00:00", "нет")],
                                tmp_path / "journal.csv")
    text = Path(out).read_text(encoding="utf-8-sig").splitlines()
    assert text[0].startswith("Группа;Фамилия") and text[1] == "TI-235;Anghis;Dmitrii;7;21.10.2026 10:00:00;нет"


# --- переназначение и массовая выдача ---

TICKETS = {1: ["a", "b", "c"], 2: ["d", "e", "f"], 3: ["g", "h", "i"]}


def test_reassign_random_differs_and_becomes_current(tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 1, False)
    for seed in range(10):
        current = td.find_first_ticket(results, "TI-235", "Anghis", "Dmitrii")
        new = td.reassign_ticket(results, TICKETS, "TI-235", "Anghis", "Dmitrii",
                                 rng=random.Random(seed))
        assert new != current
        assert td.find_first_ticket(results, "TI-235", "Anghis", "Dmitrii") == new
    assert rows(results)[-1][5] == td.REASSIGNED


def test_repeat_after_reassign_gives_new_ticket(tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 1, False)
    td.reassign_ticket(results, TICKETS, "TI-235", "Anghis", "Dmitrii", number=3)
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 3, True)
    assert td.assign_ticket(results, TICKETS, "TI-235", "Anghis", "Dmitrii") == (3, True)


def test_reassign_manual_number_validation(tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 2, False)
    for bad in ("x", "7", 2):
        with pytest.raises(td.TicketError):
            td.reassign_ticket(results, TICKETS, "TI-235", "Anghis", "Dmitrii", number=bad)
    assert td.reassign_ticket(results, TICKETS, "TI-235", "Anghis", "Dmitrii", number=" 3 ") == 3
    assert len(rows(results)) == 3


def test_reassign_without_tickets(tmp_path):
    with pytest.raises(td.TicketError):
        td.reassign_ticket(tmp_path / "r.xlsx", {}, "TI-235", "Anghis", "Dmitrii")


def test_assign_group_skips_students_with_ticket(tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 1, False)
    issued, skipped = td.assign_group(results, TICKETS, "TI-235",
                                      [("Anghis", "Dmitrii"), ("Popescu", "Ion"), ("Rusu", "Ana")],
                                      rng=random.Random(1))
    assert skipped == 1
    assert [(l, f) for l, f, _ in issued] == [("Popescu", "Ion"), ("Rusu", "Ana")]
    assert all(n in TICKETS for _, _, n in issued)
    assert len(rows(results)) == 4
    assert td.assign_group(results, TICKETS, "TI-235", [("Rusu", "Ana")]) == ([], 1)


def test_stats_count_reassigned(tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "G", "A", "B", 1, False)
    td.append_result(results, "G", "A", "B", 2, td.REASSIGNED)
    td.append_result(results, "G", "A", "B", 2, True)
    stats = td.results_stats(td.load_results(results))
    assert (stats["issued"], stats["repeats"], stats["reassigned"]) == (3, 1, 1)


# --- правка студентов и групп ---

def test_update_student_rename_updates_journal(students, tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Popescu", "Ion", 2, False)
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 1, False)
    assert td.update_student(students, "TI-235", ("Popescu", "Ion"), "TI-235", " Popescu ", "Ionel",
                             results) == ("TI-235", "Popescu", "Ionel")
    assert td.load_groups(students)["TI-235"] == [("Anghis", "Dmitrii"), ("Popescu", "Ionel")]
    assert rows(results)[1][:3] == ["TI-235", "Popescu", "Ionel"]
    assert rows(results)[2][:3] == ["TI-235", "Anghis", "Dmitrii"]


def test_update_student_move_to_group_keeps_ticket(students, tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 3, False)
    td.update_student(students, "TI-235", ("Anghis", "Dmitrii"), "TI-237", "Anghis", "Dmitrii",
                      results)
    groups = td.load_groups(students)
    assert groups["TI-235"] == [("Popescu", "Ion")]
    assert groups["TI-237"] == [("Anghis", "Dmitrii")]
    assert td.find_first_ticket(results, "TI-237", "Anghis", "Dmitrii") == 3
    assert td.find_first_ticket(results, "TI-235", "Anghis", "Dmitrii") is None


def test_update_student_rejects_duplicate_and_bad_name(students):
    with pytest.raises(td.DuplicateStudentError):
        td.update_student(students, "TI-235", ("Popescu", "Ion"), "TI-235", "Anghis", "Dmitrii")
    with pytest.raises(td.InvalidNameError):
        td.update_student(students, "TI-235", ("Popescu", "Ion"), "TI-235", "Pop3scu", "Ion")
    with pytest.raises(td.SourceError):
        td.update_student(students, "TI-235", ("Nobody", "X"), "TI-235", "Nobody", "X")
    # регистр своей же фамилии менять можно
    td.update_student(students, "TI-235", ("Popescu", "Ion"), "TI-235", "Popescu", "Ion")
    assert td.load_groups(students)["TI-235"][1] == ("Popescu", "Ion")


def test_rename_group_updates_journal(students, tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "TI-235", "Anghis", "Dmitrii", 1, False)
    assert td.rename_group(students, "TI-235", "TI-235A", results) == "TI-235A"
    assert list(td.load_groups(students)) == ["TI-235A", "TI-237"]
    assert rows(results)[1][0] == "TI-235A"
    with pytest.raises(td.DuplicateGroupError):
        td.rename_group(students, "TI-235A", "TI-237")
    with pytest.raises(td.InvalidNameError):
        td.rename_group(students, "TI-235A", "ti-238")


def test_add_students_bulk_reports_bad_lines(students):
    added, errors = td.add_students_bulk(students, "TI-237",
                                         "Rusu Ana\n\nbadname\nCebotari 2Ion\nMoraru Maria\nRusu Ana\nlupu Ion\n")
    assert added == [("Rusu", "Ana"), ("Moraru", "Maria")]
    assert [e.split(":")[0] for e in errors] == ["строка 3", "строка 4", "строка 6", "строка 7"]
    assert td.load_groups(students)["TI-237"] == [("Rusu", "Ana"), ("Moraru", "Maria")]


def test_remove_result_checks_row_values(tmp_path):
    results = tmp_path / "r.xlsx"
    td.append_result(results, "G", "A", "B", 1, False)
    td.append_result(results, "G", "C", "D", 2, False)
    loaded = td.load_results(results)
    assert td.remove_result(results, 1, loaded[0]) is False
    assert td.remove_result(results, 5, loaded[0]) is False
    assert td.remove_result(results, 0, loaded[0]) is True
    assert td.load_results(results) == [loaded[1]]
